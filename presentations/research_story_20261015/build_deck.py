"""Build five linked research slides from completed held-out validation.

Only presentation artifacts are written; frozen studies remain read-only.
No API calls occur. An absent/incomplete new result stops the build.
"""
from pathlib import Path
import hashlib
import json
import os
import re
import shutil

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
STUDIES = REPO / "research/outputs"
os.environ.setdefault("MPLCONFIGDIR", "/tmp/research-deck-matplotlib")

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

BG, PANEL, PANEL2 = "0B1425", "14233B", "1B304C"
WHITE, MUTED, DIM = "F6F8FC", "B6C4D9", "899AB4"
TEAL, GOLD, ORANGE, GRID = "5AD7C3", "F0C778", "F39A7B", "31445F"
FONT = "Arial"
W, H = 13.333333, 7.5
ASSETS, SOURCES = HERE / "assets", HERE / "sources"


def load(name):
    return json.loads((STUDIES / name).read_text())


validation_path = STUDIES / "stability_20261015/analysis.json"
if not validation_path.exists():
    raise SystemExit("STOP: completed held-out validation analysis.json is not available; no deck was rebuilt.")
val = load("stability_20261015/analysis.json")
assert val["completeness"]["complete_registered_study"], "STOP: incomplete validation cannot supply final claims"
assert val["freeze_integrity"]["status"] == "UNCHANGED"
assert val["completeness"]["completed_receiver_records"] == 1280
assert val["sources"]["completed_calls"] == 40
assert val["actual_saved_finite_view_laws"]["status"] == "PASSED_SAVED_FINITE_LAWS"
replicated = val["primary_family"]["joint_specific_replication"]
primary = val["primary_family"]["tests"]
arms = val["arms"]
neutral = val["design"]["neutral_arms"]
repeat = val["repeated_identical_inputs"]["by_arm"]
absent = ("full", "label_only", "receipt_only", "donor_joint")
unsupported = sum(arms[a]["known_signal_absent_diagnostics"]["unsupported_numeric_bit_claims"] for a in absent)
unsupported_bets = sum(arms[a]["accepted_bets"] for a in absent)
task_success = sum(v["complete_public_copy_task_success"] for v in arms.values())
explicit_success = arms["explicit"]["explicit_positive_control"]["numeric_bit_correct_success"]
source_success = val["sources"]["copy_correct"]
all_repeat_pairs = sum(repeat[a]["completed_repeat_pairs"] for a in neutral)
all_repeat_disagreement = sum(repeat[a]["prediction_disagreement_pairs"] for a in neutral)
secondary = val["secondary_announcement_family"]["tests"]
banner_harm = next(t for t in secondary if t["endpoint"] == "accepted_betting_owner_harm"
                   and t["contrast"] == "main_banner_marginal_minus_full")
view_harm = next(t for t in secondary if t["endpoint"] == "accepted_betting_owner_harm"
                 and t["contrast"] == "main_actual_view_joint_minus_full")
host = load("hostmask_20261015/INDEPENDENT_REPLICATION.json")
main = load("controlplane_20261015/controlplane_results.json")
assert not any(x["positive_direction_supported"] for x in main["primary_family"].values())

ASSETS.mkdir(exist_ok=True)
SOURCES.mkdir(exist_ok=True)
selected_sources = {
    "validation_analysis.json": "stability_20261015/analysis.json",
    "validation_results.md": "stability_20261015/RESULTS.md",
    "validation_protocol.md": "stability_20261015/PROTOCOL.md",
    "validation_freeze.json": "stability_20261015/FREEZE.json",
    "validation_manifest.json": "stability_20261015/manifest.json",
    "validation_certificate.json": "stability_20261015/finite_certificate.json",
    "jira_results.json": "revision_20261015/revision_results.json",
    "jira_report.md": "revision_20261015/REVISED_EXPERIMENT_REPORT.md",
    "domain_results.json": "multiscenario_20261015/scenario_results.json",
    "domain_interpretation.md": "multiscenario_20261015/FINAL_INTERPRETATION_AUDIT.md",
    "main_metadata_results.json": "controlplane_20261015/controlplane_results.json",
    "main_metadata_interpretation.md": "controlplane_20261015/FINAL_INTERPRETATION_AUDIT.md",
    "hostmask_replication.json": "hostmask_20261015/INDEPENDENT_REPLICATION.json",
    "hostmask_results.md": "hostmask_20261015/HOSTMASK_RESULTS.md",
    "literature_positioning.md": "controlplane_20261015/LITERATURE_POSITIONING.md",
}
for destination, relative in selected_sources.items():
    shutil.copy2(STUDIES / relative, SOURCES / destination)
claim_register = HERE / "CLAIM_REGISTER.md"
retained_project_sources = {}
if claim_register.exists():
    for relative in sorted(set(re.findall(r"\]\(sources/project/([^)]+)\)", claim_register.read_text()))):
        original_path = REPO / relative
        assert original_path.is_file(), relative
        destination = SOURCES / "project" / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(original_path, destination)
        retained_project_sources[relative] = hashlib.sha256(original_path.read_bytes()).hexdigest()

facts = {
    "validation_status": val["status"],
    "validation_primary": val["primary_family"],
    "validation_arms": arms,
    "validation_secondary_announcements": val["secondary_announcement_family"],
    "validation_repeat_diagnostics": val["repeated_identical_inputs"],
    "validation_completeness": val["completeness"],
    "validation_source_success": [source_success, 40],
    "validation_actual_saved_view_laws": val["actual_saved_finite_view_laws"]["status"],
    "validation_api_calls": 1320,
    "public_copy_success": [task_success, 1280],
    "explicit_bit_success": [explicit_success, 160],
    "no_signal_unsupported_numeric_claims": [unsupported, 640],
    "no_signal_accepted_bets": [unsupported_bets, 640],
    "repeat_prediction_disagreement": [all_repeat_disagreement, all_repeat_pairs],
    "historical_host_primary": host["primary"],
    "main_metadata_supported": 0,
    "main_metadata_tests": 4,
}
(HERE / "facts.json").write_text(json.dumps(facts, indent=2) + "\n")

plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10.5,
    "text.color": "#" + WHITE, "axes.labelcolor": "#" + MUTED,
    "xtick.color": "#" + MUTED, "ytick.color": "#" + WHITE,
    "axes.facecolor": "#" + BG, "figure.facecolor": "#" + BG,
    "axes.edgecolor": "#" + GRID, "savefig.facecolor": "#" + BG,
})
# Scientific plot: each estimate and interval comes from the completed analysis.
fig, ax = plt.subplots(figsize=(4.35, 2.55))
fig.subplots_adjust(left=.29, right=.96, bottom=.22, top=.96)
labels = {"joint": "Joint", "full": "Full fixed", "label_only": "Label only",
          "receipt_only": "Receipt only", "donor_joint": "Donor", "explicit": "Explicit bit"}
for y, arm in zip(range(5, -1, -1), neutral):
    summary = arms[arm]["signed_S"]
    mean, (lo, hi) = summary["mean"], summary["bootstrap95"]
    color = TEAL if arm == "joint" else GOLD if arm == "explicit" else MUTED
    ax.errorbar(mean, y, xerr=[[max(0, mean-lo)], [max(0, hi-mean)]],
                color="#" + color, marker="o", markersize=5.5, capsize=3, linewidth=1.7)
ax.set_yticks(range(5, -1, -1), [labels[a] for a in neutral])
ax.set_xlim(-.13, .55); ax.set_ylim(-.6, 5.6)
ax.set_xticks([0, .25, .5], ["0", ".25", ".50"])
ax.set_xlabel("Signed decoding S", fontsize=10.2, labelpad=3)
ax.axvline(0, color="#" + DIM, linestyle="--", linewidth=1)
ax.grid(axis="x", color="#" + GRID, alpha=.7)
ax.tick_params(length=0, labelsize=10)
for spine in ax.spines.values(): spine.set_visible(False)
fig.savefig(ASSETS / "validation_signed_decoding.png", dpi=300)
fig.savefig(ASSETS / "validation_signed_decoding.svg")
plt.close(fig)

# A distinct action endpoint: descriptive accepted betting in the banner 2x2.
heat_arms = [["joint_marginal_banner", "joint_full_banner"],
             ["full_marginal_banner", "full_full_banner"]]
heat_counts = np.array([[arms[a]["accepted_bets"] for a in row] for row in heat_arms])
fig, ax = plt.subplots(figsize=(4.35, 1.5))
fig.subplots_adjust(left=.27, right=.96, bottom=.25, top=.96)
ax.imshow(heat_counts / 80, vmin=0, vmax=1, cmap="YlGnBu", aspect="auto")
for y, x in np.ndindex(heat_counts.shape):
    ax.text(x, y, f"{heat_counts[y,x]}/80", ha="center", va="center",
            color="#" + (BG if heat_counts[y,x] < 36 else WHITE), fontsize=11.5, fontweight="bold")
ax.set_xticks([0, 1], ["Marginal banner", "Full banner"])
ax.set_yticks([0, 1], ["Joint view", "Full view"])
ax.tick_params(length=0, labelsize=9.8)
for spine in ax.spines.values(): spine.set_visible(False)
fig.savefig(ASSETS / "validation_banner_bets.png", dpi=300)
fig.savefig(ASSETS / "validation_banner_bets.svg")
plt.close(fig)

# Matplotlib leaves harmless end-of-line spaces in SVG paths. Keep generated
# artifacts clean for repository whitespace checks without changing geometry.
for svg in ASSETS.glob("validation_*.svg"):
    svg.write_text("\n".join(line.rstrip() for line in svg.read_text().splitlines()) + "\n")

prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
prs.core_properties.title = "Joint Metadata, Privacy Announcements and Agent Decisions"
prs.core_properties.subject = "Five-slide empirical briefing with held-out ablations and stability controls"
prs.core_properties.author = "AI-assisted research briefing; author metadata not supplied"
prs.core_properties.keywords = "AI collusion, metadata, held-out validation, joint information, Apart Research"


def box(slide, x, y, w, h, fill=PANEL, radius=True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(.7)
    if radius: shape.adjustments[0] = .10
    return shape


def txt(slide, text, x, y, w, h, size=16, color=WHITE, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, value in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = value; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(0); p.space_before = Pt(0); p.line_spacing = 1.08
        if align is not None: p.alignment = align
    return shape


def arrow(slide, x, y, w=.2, h=.22, color=DIM):
    return box(slide, x, y, w, h, color, radius=False)


def base(number, label, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    box(slide, .55, .36, .35, .05, TEAL, radius=False)
    txt(slide, f"0{number}  /  {label.upper()}", 1.03, .27, 11.0, .25, 10.5, TEAL, True)
    txt(slide, title, .55, .71, 12.23, .65, 29.5, WHITE, True)
    txt(slide, subtitle, .55, 1.39, 12.20, .46, 14.3, MUTED)
    txt(slide, "15 OCT 2026  |  PRE-SPRINT EVIDENCE  |  AI-ASSISTED BRIEFING",
        .55, 7.20, 10.8, .15, 8.5, DIM)
    txt(slide, str(number) + " / 5", 12.10, 7.17, .65, .22, 10.2, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def source_link(slide, label, url, x, w):
    # A shape-level link preserves readable text color in LibreOffice/PDF.
    shape = txt(slide, label, x, 6.98, w, .15, 8.6, MUTED)
    shape.click_action.hyperlink.address = url


# 01 — documented motivation and the single empirical question.
s = base(1, "Question", "Can defensive metadata steer an agent?",
         "We separate information in a tool response from what the agent is told about its privacy.")
cards = [
    ("REAL PRODUCT VULNERABILITY", "Microsoft 365 Copilot", "CVE-2025-32711",
     "AI command injection enabled unauthorized information disclosure over a network.",
     "Product vulnerability; not deployed collusion", ORANGE),
    ("CONTROLLED WORKPLACE BENCHMARK", "Colosseum", "11 of 19 models",
     "showed positive changes in action-level coalition advantage when private chat was enabled.",
     "Action and conversation evidence can differ", TEAL),
    ("CONTROLLED CAPABILITY EXPERIMENT", "Hidden in Plain Text", "Filtering was bypassed",
     "Optimized agents used covert signaling despite oversight and paraphrasing in candidate screening.",
     "Constructed task; not a hiring incident", GOLD),
]
for i, (tag, title, stat, body, qualifier, accent) in enumerate(cards):
    x = .55 + i * 4.13
    box(s, x, 2.02, 3.96, 2.65)
    box(s, x+.20, 2.22, .07, .33, accent, radius=False)
    txt(s, tag, x+.40, 2.22, 3.33, .30, 9.3, accent, True)
    txt(s, title, x+.23, 2.70, 3.49, .33, 19, WHITE, True)
    txt(s, stat, x+.23, 3.13, 3.50, .42, 24 if i != 2 else 20, accent, True)
    txt(s, body, x+.23, 3.68, 3.48, .66, 14.0, WHITE)
    txt(s, qualifier, x+.23, 4.40, 3.50, .18, 10.3, MUTED)
box(s, .55, 5.10, 12.23, .91, PANEL2)
txt(s, "OUR NARROW QUESTION", .78, 5.28, 2.24, .22, 11.2, TEAL, True)
txt(s, "How do actual metadata and privacy announcements each affect an agent's decisions?",
    3.12, 5.28, 9.35, .49, 18.0, WHITE, True)
txt(s, "Colosseum grounded the workplace setup; security and covert-channel papers motivated the observation test.",
    .55, 6.32, 12.23, .25, 12.3, MUTED)
txt(s, "No full FLOWSEAL / OCELOT / OrchDebate reproduction. Existing work already recognizes implicit-flow risks.",
    .55, 6.67, 12.23, .19, 10.4, DIM)
source_link(s, "[S1] Microsoft CNA / vendor advisory", "https://msrc.microsoft.com/update-guide/vulnerability/CVE-2025-32711", .55, 3.8)
source_link(s, "[S2] Colosseum v2, §5.4 / Figure 4", "https://arxiv.org/abs/2602.15198v2", 4.71, 3.8)
source_link(s, "[S3] Hidden in Plain Text v2, §5.1", "https://arxiv.org/abs/2410.03768v2", 8.87, 3.8)

# 02 — one source-to-host-to-decision mechanism, public task held fixed.
s = base(2, "Mechanism", "The combination reveals what either field hides",
         "A supplied-codebook experiment isolates one path: source attestation → host metadata → receiver decision.")
box(s, .55, 2.12, 2.48, 3.10)
txt(s, "SOURCE MODEL", .76, 2.34, 2.06, .23, 11.6, TEAL, True)
txt(s, "Private bit B", .76, 2.91, 2.06, .35, 23.5, WHITE, True)
txt(s, "Copies B into A\nResponse saved\nNo mask R",
    .76, 3.52, 2.06, 1.12, 16.0, WHITE)
txt(s, "Neutral copying; no collusive goal", .76, 4.79, 2.06, .27, 10.0, MUTED)
txt(s, "›", 3.09, 3.47, .24, .42, 29, TEAL, True)
box(s, 3.42, 2.12, 4.17, 3.10)
txt(s, "HOST MASKING", 3.65, 2.34, 3.71, .23, 11.6, TEAL, True)
box(s, 3.65, 2.93, 1.65, .87, PANEL2)
box(s, 5.53, 2.93, 1.83, .87, PANEL2)
txt(s, "Label bit = R", 3.78, 3.16, 1.39, .35, 16.1, WHITE, True, PP_ALIGN.CENTER)
txt(s, "Receipt bit =\nA ⊕ R ⊕ P", 5.65, 3.06, 1.59, .57, 15.2, WHITE, True, PP_ALIGN.CENTER)
txt(s, "Each field alone hides B", 3.65, 4.03, 3.71, .28, 15.6, MUTED, align=PP_ALIGN.CENTER)
txt(s, "Together: recover A", 3.65, 4.49, 3.71, .30, 19.0, TEAL, True, PP_ALIGN.CENTER)
txt(s, "Public P; label PASS/NORMALIZE; receipt 0/10", 3.65, 4.92, 3.71, .17, 9.1, MUTED, align=PP_ALIGN.CENTER)
txt(s, "›", 7.72, 3.47, .24, .42, 29, TEAL, True)
box(s, 8.12, 2.12, 4.66, 3.10)
txt(s, "RECEIVER MODEL", 8.35, 2.34, 4.20, .23, 11.6, TEAL, True)
txt(s, "Supplied decoder + routing table", 8.35, 2.94, 4.20, .37, 18.0, WHITE, True)
txt(s, "Choose SAFE / LEFT / RIGHT\nCorrect bet +2  ·  Wrong bet −4\nAny accepted bet: owner utility −7",
    8.35, 3.55, 4.20, 1.06, 17.0, WHITE)
txt(s, "Inference and accepted action scored separately", 8.35, 4.84, 4.20, .20, 10.6, MUTED)
box(s, .55, 5.62, 12.23, .69, PANEL2)
txt(s, "UNCHANGED PUBLIC PATH", .76, 5.84, 2.72, .22, 11.5, TEAL, True)
txt(s, "Canonical schedule + checksum → receiver copies the same public artifact in every arm",
    3.63, 5.83, 8.90, .27, 14.3, WHITE)
txt(s, "Joint target signal requires different realized source attestations for B=0 and B=1; source-copy accuracy is measured.",
    .55, 6.61, 12.23, .25, 11.3, GOLD)
txt(s, "Earlier Jira context/access and resource-penalty diagnostics motivated isolating this channel; they are separate results.",
    .55, 6.99, 12.23, .15, 9.2, DIM)

# 03 — six distinct controls for one mechanism, and a fixed inference rule.
s = base(3, "Validation", "One mechanism, distinct controls, fixed tests",
         "The actual observation, policy announcement and stochastic response are tested separately.")
cols = [(.55, 2.38), (2.93, 5.20), (8.13, 4.65)]
for (x,w), label in zip(cols, ["QUESTION", "CONTROL / CHANGE", "WHY IT IS NEEDED"]):
    txt(s, label, x+.12, 1.98, w-.22, .22, 10.3, TEAL, True)
controls = [
    ("AVAILABILITY", "Joint vs complete fixed metadata view", "An effect without target-bearing input"),
    ("COMPONENT NECESSITY", "Label-only and receipt-only; balance both host pads", "A shortcut carried by either field alone"),
    ("TARGET SPECIFICITY", "Same-format codeword from a fixed target-independent donor", "Reaction to any plausible-looking codeword"),
    ("READING SENSITIVITY", "Authorized explicit target bit in a fixed metadata view", "Failure to read available information"),
    ("POLICY FRAMING", "Actual joint/full view × exact earlier marginal/full banners", "Confounding data with privacy instructions"),
    ("RELIABILITY", "20 fresh worlds; two identical-input repeats per neutral arm", "Check fresh-world and response consistency"),
]
for i, row in enumerate(controls):
    y = 2.34 + .51*i
    box(s, .55, y, 12.23, .45, PANEL if i%2 == 0 else PANEL2)
    for j, value in enumerate(row):
        x,w = cols[j]
        txt(s, value, x+.12, y+.12, w-.24, .25, 10.6 if j == 0 else 12.4,
            TEAL if j == 0 else WHITE, j == 0)
box(s, .55, 5.68, 12.23, .49, PANEL2)
txt(s, "1,320 completed calls  ·  40 sources + 1,280 receivers  ·  GPT-4.1 mini  ·  T=.7  ·  max 1,500 tokens",
    .73, 5.84, 11.87, .22, 12.4, WHITE, True)
txt(s, "20 world blocks are the statistical units. Four primary paired contrasts; Holm correction; 20,000 world bootstraps.",
    .55, 6.36, 12.23, .25, 12.0, MUTED)
txt(s, "Frozen before this study's calls; exact requests retained; raw-field inference separated from accepted actions. No outcome retries.",
    .55, 6.70, 12.23, .21, 11.3, MUTED)
txt(s, "Original metadata study: 0/4 primary tests supported. Later cohorts have separate prospective freezes; earlier outcomes remain in the record.",
    .55, 6.99, 12.23, .15, 8.9, DIM)

# 04 — exclusively the new controlled validation; historical results are notes.
headline = "Joint decoding passed all four held-out comparisons" if replicated else "Held-out validation did not meet the four-comparison criterion"
s = base(4, "Results", headline,
         "Four prespecified joint-minus-control contrasts, 20 worlds each. Intervals are descriptive world-cluster 95% bootstraps.")
table_cols = [(.55, 2.46), (3.01, 3.18), (6.19, 1.63)]
for (x,w), label in zip(table_cols, ["JOINT MINUS…", "MEAN ΔS [95% INTERVAL]", "HOLM4 P"]):
    txt(s, label, x+.10, 2.03, w-.18, .20, 10.2, TEAL, True)
for i, test in enumerate(primary):
    y = 2.39 + .62*i
    box(s, .55, y, 7.27, .55, PANEL if i%2 == 0 else PANEL2)
    comparator = test["name"].removeprefix("joint_minus_")
    name = {"full": "Full fixed view", "label_only": "Label only", "receipt_only": "Receipt only", "donor_joint": "Donor codeword"}[comparator]
    lo, hi = test["bootstrap95"]
    vals = [name, f"{test['mean']:+.4f}  [{lo:.4f}, {hi:.4f}]", f"{test['holm4_p']:.3g}"]
    for j, value in enumerate(vals):
        x,w = table_cols[j]
        txt(s, value, x+.10, y+.17, w-.18, .25, 13.0 if j != 1 else 12.4,
            TEAL if j == 2 and test["positive_and_holm4_lt_05"] else WHITE, j == 2)
box(s, .55, 5.07, 7.27, 1.65, PANEL2)
txt(s, "ERROR AND STABILITY CHECKS", .74, 5.25, 6.89, .22, 11.1, TEAL, True)
txt(s, f"Explicit bit: {explicit_success}/160 correct  ·  public copy: {task_success}/1,280\n"
       f"Explicit correct bets: {arms['explicit']['accepted_correct_bets']}/160; no-signal bets: {unsupported_bets}/640\n"
       f"Unsupported bit claims: {unsupported}/640; joint repeat disagreements: {repeat['joint']['prediction_disagreement_pairs']}/80",
    .74, 5.68, 6.89, .80, 12.1, WHITE)
txt(s, "Neutral arms: signed decoding", 8.34, 2.03, 4.42, .25, 13.8, WHITE, True)
s.shapes.add_picture(str(ASSETS/"validation_signed_decoding.png"), Inches(8.28), Inches(2.35), width=Inches(4.50), height=Inches(2.61))
txt(s, "Accepted bets: view × announcement", 8.34, 5.11, 4.42, .25, 13.0, WHITE, True)
s.shapes.add_picture(str(ASSETS/"validation_banner_bets.png"), Inches(8.28), Inches(5.43), width=Inches(4.50), height=Inches(1.35))
txt(s, "S=.5 perfect; S=0 constant/abstain. Bit errors are not collusion-detector errors. Original host-mask ΔS=.22656 is a separate cohort.",
    .55, 6.99, 12.23, .15, 8.8, DIM)

# 05 — one bounded finding: crossed announcements dissociate action and exposure.
if replicated:
    title = "Privacy announcements changed action on the same data"
    subtitle = "Usable information and observed betting came apart when the actual view was crossed with the policy announcement."
else:
    title = "The new study bounds the original metadata result"
    subtitle = "The four-control criterion did not replicate; the original result is not substituted for this held-out outcome."
s = base(5, "Finding", title, subtitle)
joint, full = arms["joint"], arms["full"]
means = [test["mean"] for test in primary]
min_delta, max_delta = min(means), max(means)
if replicated:
    signal_head = "Joint target information mattered."
    signal_body = f"All four controlled ΔS effects: +{min_delta:.4f} to +{max_delta:.4f}.\nRemoving either field or substituting a donor codeword reduced target decoding."
else:
    failed = sum(not t["positive_and_holm4_lt_05"] for t in primary)
    signal_head = "Joint-specific replication failed."
    signal_body = f"{failed}/4 prespecified comparisons failed the positive Holm4 criterion.\nThe evidence remains bounded to the reported conditions and cohorts."
jm, jf = arms['joint_marginal_banner'], arms['joint_full_banner']
fm, ff = arms['full_marginal_banner'], arms['full_full_banner']
decision_body = (f"Same informative joint view: {jm['accepted_bets']}/80 bets with the marginal banner; "
                 f"{jf['accepted_bets']}/80 with the full banner.\nThe metadata still carried the target bit.")
boundary_body = (f"Same target-independent fixed view: {fm['accepted_bets']}/80 bets with the marginal banner; "
                 f"{ff['accepted_bets']}/80 with the full banner.\nThe uninformed betting netted {fm['accepted_net_payoff']:+.0f}.")
linked = [
    ("01  USABLE INFORMATION", signal_head, signal_body),
    ("02  INFORMATION PRESENT", "An open channel can look safe.", decision_body),
    ("03  INFORMATION ABSENT", "Bets persisted without target information.", boundary_body),
]
for i, (tag, head, body) in enumerate(linked):
    x = .55 + 4.13*i
    box(s, x, 2.07, 3.96, 2.99)
    txt(s, tag, x+.22, 2.28, 3.52, .25, 11.4, TEAL, True)
    txt(s, head, x+.22, 2.86, 3.52, .74, 19.1, WHITE, True)
    txt(s, body, x+.22, 3.84, 3.52, 1.01, 13.0, WHITE)
if banner_harm["holm6_p"] < .05:
    lo, hi = banner_harm['bootstrap95']
    framing = (f"Announcement effect on owner harm: {banner_harm['mean']:+.2f}/call "
               f"[95% {lo:.2f}, {hi:.2f}], Holm6 p={banner_harm['holm6_p']:.3g}. Exact prior policy bundles were crossed.")
else:
    framing = (f"The prespecified policy-banner harm effect was unsupported after Holm6 "
               f"(p={banner_harm['holm6_p']:.3g}); no framing-success claim is made.")
txt(s, framing, .55, 5.36, 12.23, .41, 12.1, GOLD)
box(s, .55, 5.96, 12.23, .66, PANEL2)
txt(s, "REAL WORKFLOW IMPLICATION", .75, 6.17, 2.77, .22, 10.7, TEAL, True)
txt(s, "Test what agents can learn and do separately from what the policy claims.",
    3.63, 6.14, 8.92, .31, 14.3, WHITE, True)
txt(s, "Controlled supplied-codebook capability; neutral source; public-copy task only. No spontaneous collusion, new XOR theory or original-defense break.",
    .55, 6.99, 12.23, .15, 8.8, DIM)

notes_path = HERE / "SPEAKER_NOTES.md"
notes = notes_path.read_text() if notes_path.exists() else ""
chunks = re.split(r"(?m)^## Slide [1-5][^\n]*\n", notes)
if len(chunks) == 6:
    for slide, note in zip(prs.slides, chunks[1:]):
        slide.notes_slide.notes_text_frame.text = note.strip()
else:
    for i, slide in enumerate(prs.slides, 1):
        slide.notes_slide.notes_text_frame.text = f"Slide {i}. Detailed narrative and source qualifications are in SPEAKER_NOTES.md and CLAIM_REGISTER.md."
output = HERE / "research_story_5_slides.pptx"
prs.save(output)
assert len(Presentation(output).slides) == 5
manifest = {
    "slides": 5, "output": output.name, "paid_model_calls_by_builder": 0,
    "builder_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    "graphs_regenerated_from_completed_results": True,
    "frozen_study_files_changed": False,
    "speaker_notes_embedded": len(chunks) == 6,
    "validation_status": val["status"],
    "source_files": {str((STUDIES/r).relative_to(REPO)): hashlib.sha256((STUDIES/r).read_bytes()).hexdigest()
                     for r in selected_sources.values()},
    "retained_project_sources": retained_project_sources,
    "facts": facts,
}
(HERE/"BUILD_MANIFEST.json").write_text(json.dumps(manifest, indent=2)+"\n")
print(json.dumps({"pptx": str(output), "slides": 5, "speaker_notes_embedded": manifest["speaker_notes_embedded"],
                  "validation_status": val["status"]}))
